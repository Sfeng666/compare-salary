"""Rebuild attributed public data. Pass a directory containing publisher downloads."""
import hashlib, json, sys, zipfile, csv
from pathlib import Path
from datetime import datetime, timezone
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
SOURCE=Path(sys.argv[1]) if len(sys.argv)>1 else ROOT/'.local/sources'
def write(name,value):
    (ROOT/'data'/name).write_text(json.dumps(value,separators=(',',':')))
with (SOURCE/'census-ct-crosswalk.txt').open(encoding='utf-8-sig',newline='') as f:
    ct=list(csv.reader(f,delimiter='|'))[1:]
ct_counties={r[7]:r[0]+r[1] for r in ct if len(r)>8 and r[5]!='00000'}
ct_towns={r[8].lower().removesuffix(' town').removesuffix(' city'):r[0]+r[1] for r in ct if len(r)>8 and r[5]!='00000'}
counties={};states={}
epi=openpyxl.load_workbook(SOURCE/'epi-2026.xlsx',read_only=True,data_only=True)
for r in epi['County'].iter_rows(min_row=3,values_only=True):
    if r[4] not in ('1p0c','2p0c'):continue
    fips=str(int(r[2])).zfill(5);states[r[1]]=fips[:2]
    c=counties.setdefault(fips,{'fips':fips,'state':r[1],'name':r[3],'families':{},'rents':[]})
    c['families'][r[4]]=dict(zip(['food','transport','healthcare','other'],r[14:18]))
hud=openpyxl.load_workbook(SOURCE/'hud-2027.xlsx',read_only=True,data_only=True).worksheets[0]
for r in hud.iter_rows(min_row=2,values_only=True):
    if not r[0]:break
    fips=str(r[7]).zfill(10)[:5]
    if r[0]=='CT':fips=ct_counties.get(str(r[7]).zfill(10),fips)
    if fips in counties:counties[fips]['rents'].append({'area':r[6],'town':r[4] or '', 'population':r[8] or 1,'oneBedroom':r[10],'twoBedroom':r[11],'metro':str(r[2])[5:10] if str(r[2]).startswith('METRO') else None})
cities={}
with zipfile.ZipFile(SOURCE/'geonames-US.zip') as z:
    for line in z.read('US.txt').decode().splitlines():
        r=line.split('\t')
        if r[6]!='P' or r[7] not in ('PPL','PPLX','PPLA','PPLA2','PPLA3','PPLA4','PPLC') or r[10] not in states:continue
        fips=states[r[10]]+r[11].zfill(3)
        if fips not in counties and r[10]=='CT':fips=ct_counties.get('09'+r[11].zfill(3)+r[12].zfill(5),ct_towns.get(r[1].lower(),fips))
        if fips not in counties:continue
        for name in set([r[1],r[2]]):
            key=f'{name}, {r[10]}'.lower();row={'city':name,'state':r[10],'fips':fips,'population':int(r[14] or 0)}
            if key not in cities or row['population']>cities[key]['population']:cities[key]=row
wages={}
for archive,suffix,is_state in [('bls-2025-metro.zip','MSA_M2025_dl.xlsx',False),('bls-2025-state.zip','state_M2025_dl.xlsx',True)]:
    with zipfile.ZipFile(SOURCE/archive) as z:
        sheet=openpyxl.load_workbook(z.open(next(n for n in z.namelist() if n.endswith(suffix))),read_only=True,data_only=True).active
        it=sheet.iter_rows(values_only=True);headers=next(it)
        for values in it:
            r=dict(zip(headers,values));code=r['OCC_CODE']
            if r['NAICS']!='000000' or str(r['OWN_CODE'])!='1235' or code not in ('21-1021','21-1022','21-1023','21-1029') or not isinstance(r['A_MEDIAN'],(int,float)):continue
            area='state:'+r['PRIM_STATE'] if is_state else str(r['AREA']).zfill(5)
            wages.setdefault(area,{'name':r['AREA_TITLE'],'occupations':{}})['occupations'][code]={'median':r['A_MEDIAN'],'p25':r['A_PCT25'] if isinstance(r['A_PCT25'],(int,float)) else None,'p75':r['A_PCT75'] if isinstance(r['A_PCT75'],(int,float)) else None,'employment':r['TOT_EMP'] if isinstance(r['TOT_EMP'],(int,float)) else None}
sources=[
 {'name':'Economic Policy Institute. 2026. Family Budget Calculator.','url':'https://files.epi.org/uploads/fbc_data_2026.xlsx','file':'epi-2026.xlsx','vintage':'2025 dollars','license':'CC-BY-NC-SA-4.0','changes':'Selected nonhousing categories for one or two adults, no children; housing replaced with HUD FY2027.'},
 {'name':'HUD FY2027 Fair Market Rents','url':'https://www.huduser.gov/portal/datasets/fmr/fmr2027/FY27_FMRs.xlsx','file':'hud-2027.xlsx','vintage':'FY2027','license':'US government public data'},
 {'name':'BLS May 2025 OEWS metropolitan wages','url':'https://www.bls.gov/oes/special-requests/oesm25ma.zip','file':'bls-2025-metro.zip','vintage':'May 2025','license':'US government public data'},
 {'name':'BLS May 2025 state wages','url':'https://www.bls.gov/oes/special-requests/oesm25st.zip','file':'bls-2025-state.zip','vintage':'May 2025','license':'US government public data'},
 {'name':'GeoNames US gazetteer','url':'https://download.geonames.org/export/dump/US.zip','file':'geonames-US.zip','license':'CC-BY-4.0','changes':'Selected populated places, normalized county lookup and removed unused fields.'},
 {'name':'US Census Connecticut crosswalk','url':'https://www2.census.gov/geo/docs/reference/ct_change/ct_cou_to_cousub_crosswalk.txt','file':'census-ct-crosswalk.txt','license':'US government public data'}]
for s in sources:s['sha256']=hashlib.sha256((SOURCE/s['file']).read_bytes()).hexdigest()
write('base.json',{'version':'epi2026-hud2027-bls2025-public-v1','builtAt':datetime.now(timezone.utc).isoformat(),'sources':sources,'counties':counties,'wages':wages})
write('cities.json',cities)
print(json.dumps({'counties':len(counties),'cities':len(cities),'wageAreas':len(wages)}))
